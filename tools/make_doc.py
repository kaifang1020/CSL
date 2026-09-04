# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, "/Users/maokaifang/Downloads/Avatar")
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from content import TITLE, SUBTITLE, AXES, RULES, LEVELS

FONT = "Calibri"
INK, MUT, TECH_C, PSY_C = "1F2933", "5A6672", "1D4ED8", "B45309"

def rgb(h): return RGBColor.from_string(h)

def style_run(r, size=11, bold=False, italic=False, color=INK, font=FONT):
    r.font.name, r.font.size, r.font.bold, r.font.italic = font, Pt(size), bold, italic
    r.font.color.rgb = rgb(color)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), font)

def para(doc, text="", size=11, bold=False, italic=False, color=INK,
         before=0, after=6, align=None, indent=0):
    p = doc.add_paragraph()
    if align: p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.18
    if indent: pf.left_indent = Inches(indent)
    if text: style_run(p.add_run(text), size, bold, italic, color)
    return p

def shade(cell, hexcolor):
    el = OxmlElement('w:shd'); el.set(qn('w:val'),'clear'); el.set(qn('w:fill'), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)

def cell_text(cell, runs, size=10.5, space=2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(space)
    p.paragraph_format.space_after  = Pt(space)
    p.paragraph_format.line_spacing = 1.14
    for t, kw in runs:
        kw = dict(kw); kw.setdefault("size", size)
        style_run(p.add_run(t), **kw)
    return p

doc = Document()
s = doc.sections[0]
s.left_margin = s.right_margin = Inches(0.8)
s.top_margin  = s.bottom_margin = Inches(0.75)
n = doc.styles['Normal']; n.font.name = FONT; n.font.size = Pt(11)
n.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

# ── 标题 ──
para(doc, TITLE, size=21, bold=True, after=3)
para(doc, SUBTITLE, size=10.5, italic=True, color=MUT, after=14)

# ── 总览图 ──
doc.add_picture("q1400_128.png", width=Inches(7.0))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
para(doc, "Figure 1 · The framework at a glance. Three timescales × two independent axes.",
     size=9, italic=True, color=MUT, align=WD_ALIGN_PARAGRAPH.CENTER, after=16)

# ── 两条分层规则 ──
para(doc, "How metrics are assigned to a level", size=14, bold=True, before=4, after=6)
for t, b in RULES[:2]:
    p = para(doc, after=8)
    style_run(p.add_run(t + " — "), size=11, bold=True)
    style_run(p.add_run(b), size=11)

# ── 两条轴 ──
para(doc, "The two axes", size=14, bold=True, before=10, after=6)
for tag, nm, gl in AXES:
    p = para(doc, after=5)
    style_run(p.add_run(f"{tag} · {nm}  "), size=11, bold=True,
              color=TECH_C if tag == "A" else PSY_C)
    style_run(p.add_run(gl), size=11, color=MUT)

doc.add_page_break()

# ── 三层详表 ──
for idx, L in enumerate(LEVELS, start=1):
    para(doc, f"{idx}. {L['key'].title()}-level Evaluation", size=16, bold=True, before=0, after=2)
    p = para(doc, after=4)
    style_run(p.add_run(f"{L['dur']}  ·  {L['scale']}"), size=10.5, bold=True, color=MUT)
    p = para(doc, after=10)
    style_run(p.add_run("Test: "), size=11, bold=True)
    style_run(p.add_run(L["test"]), size=11, italic=True)

    for axis_tag, axis_name, items, accent, band in (
        ("A", AXES[0][1], L["tech"], TECH_C, "EEF3FC"),
        ("B", AXES[1][1], L["psy"],  PSY_C,  "FDF3E4")):

        p = para(doc, before=6, after=5)
        style_run(p.add_run(f"{axis_tag} · {axis_name.title()}"), size=12, bold=True, color=accent)

        t = doc.add_table(rows=1, cols=3)
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0].cells
        for c, label in zip(hdr, ("Aspect", "What it tests", "Metrics")):
            shade(c, band); cell_text(c, [(label, dict(bold=True, size=10))])
        for name, question, metrics in items:
            row = t.add_row().cells
            cell_text(row[0], [(name, dict(bold=True))])
            cell_text(row[1], [(question, dict(color=INK))])
            cell_text(row[2], [(metrics, dict(color="3F4A55"))])
        for r_ in t.rows:
            r_.cells[0].width = Inches(1.55)
            r_.cells[1].width = Inches(2.65)
            r_.cells[2].width = Inches(2.80)
        para(doc, after=4)

    if idx < len(LEVELS): doc.add_page_break()

# ── 汇报原则 ──
para(doc, "Reporting principle", size=14, bold=True, before=12, after=6)
para(doc, RULES[2][1], size=11, after=10)

out = "Avatar_Evaluation_Temporal_Hierarchical_Framework.docx"
doc.save(out); print("saved", out)
