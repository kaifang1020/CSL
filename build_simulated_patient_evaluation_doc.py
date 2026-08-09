from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path


OUT = Path('/Users/maokaifang/Downloads/Avatar/Simulated_Patient_Evaluation_Framework_Psychology_Grounded.docx')
# Named Google-Docs-targeted override: a CJK-safe sans serif is required because
# LibreOffice does not reliably apply Word's eastAsia fallback when the primary
# font is Arial. PingFang keeps the understated Google Docs visual character.
PRIMARY_FONT = 'PingFang SC'
CJK_FONT = 'PingFang SC'


def set_run_font(run, name=PRIMARY_FONT, size=11, bold=None, italic=None, color='000000'):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), name)
    rfonts.set(qn('w:hAnsi'), name)
    rfonts.set(qn('w:cs'), name)
    rfonts.set(qn('w:eastAsia'), CJK_FONT)
    rfonts.set(qn('w:hint'), 'eastAsia')
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rpr.append(lang)
    lang.set(qn('w:val'), 'en-US')
    lang.set(qn('w:eastAsia'), 'zh-CN')


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(total))
    tblW.set(qn('w:type'), 'dxa')
    tblInd = tblPr.first_child_found_in('w:tblInd')
    if tblInd is None:
        tblInd = OxmlElement('w:tblInd')
        tblPr.append(tblInd)
    tblInd.set(qn('w:w'), '0')
    tblInd.set(qn('w:type'), 'dxa')
    layout = tblPr.first_child_found_in('w:tblLayout')
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tblPr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.first_child_found_in('w:tcW')
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(widths[idx]))
            tcW.set(qn('w:type'), 'dxa')
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = borders.find(qn(f'w:{edge}'))
        if e is None:
            e = OxmlElement(f'w:{edge}')
            borders.append(e)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), 'DADCE0')


def repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def add_hyperlink(paragraph, text, url, color='1155CC'):
    part = paragraph.part
    rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), PRIMARY_FONT)
    rFonts.set(qn('w:hAnsi'), PRIMARY_FONT)
    rFonts.set(qn('w:cs'), PRIMARY_FONT)
    rFonts.set(qn('w:eastAsia'), CJK_FONT)
    rFonts.set(qn('w:hint'), 'eastAsia')
    rPr.append(rFonts)
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'en-US')
    lang.set(qn('w:eastAsia'), 'zh-CN')
    rPr.append(lang)
    c = OxmlElement('w:color')
    c.set(qn('w:val'), color)
    rPr.append(c)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_numbering(document):
    numbering = document.part.numbering_part.element
    existing_abs = [int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))]
    existing_num = [int(x.get(qn('w:numId'))) for x in numbering.findall(qn('w:num'))]
    abs_b = max(existing_abs or [0]) + 1
    abs_n = abs_b + 1
    num_b = max(existing_num or [0]) + 1
    num_n = num_b + 1

    def add_abs(abs_id, fmt, text):
        abstract = OxmlElement('w:abstractNum')
        abstract.set(qn('w:abstractNumId'), str(abs_id))
        multi = OxmlElement('w:multiLevelType')
        multi.set(qn('w:val'), 'singleLevel')
        abstract.append(multi)
        lvl = OxmlElement('w:lvl')
        lvl.set(qn('w:ilvl'), '0')
        start = OxmlElement('w:start'); start.set(qn('w:val'), '1'); lvl.append(start)
        nf = OxmlElement('w:numFmt'); nf.set(qn('w:val'), fmt); lvl.append(nf)
        lt = OxmlElement('w:lvlText'); lt.set(qn('w:val'), text); lvl.append(lt)
        jc = OxmlElement('w:lvlJc'); jc.set(qn('w:val'), 'left'); lvl.append(jc)
        pPr = OxmlElement('w:pPr')
        tabs = OxmlElement('w:tabs')
        tab = OxmlElement('w:tab'); tab.set(qn('w:val'), 'num'); tab.set(qn('w:pos'), '720'); tabs.append(tab)
        pPr.append(tabs)
        ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '720'); ind.set(qn('w:hanging'), '360'); pPr.append(ind)
        lvl.append(pPr)
        rPr = OxmlElement('w:rPr')
        rfonts = OxmlElement('w:rFonts')
        rfonts.set(qn('w:ascii'), PRIMARY_FONT)
        rfonts.set(qn('w:hAnsi'), PRIMARY_FONT)
        rfonts.set(qn('w:cs'), PRIMARY_FONT)
        rfonts.set(qn('w:eastAsia'), CJK_FONT)
        rPr.append(rfonts)
        lvl.append(rPr)
        abstract.append(lvl)
        numbering.append(abstract)

    def add_num(num_id, abs_id):
        num = OxmlElement('w:num')
        num.set(qn('w:numId'), str(num_id))
        abstractNumId = OxmlElement('w:abstractNumId')
        abstractNumId.set(qn('w:val'), str(abs_id))
        num.append(abstractNumId)
        numbering.append(num)

    add_abs(abs_b, 'bullet', '●')
    add_abs(abs_n, 'decimal', '%1.')
    add_num(num_b, abs_b)
    add_num(num_n, abs_n)
    return num_b, num_n


def apply_num(paragraph, num_id):
    pPr = paragraph._p.get_or_add_pPr()
    numPr = pPr.find(qn('w:numPr'))
    if numPr is None:
        numPr = OxmlElement('w:numPr')
        pPr.append(numPr)
    ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), '0'); numPr.append(ilvl)
    numId = OxmlElement('w:numId'); numId.set(qn('w:val'), str(num_id)); numPr.append(numId)
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15


def add_bullet(doc, text, num_b):
    p = doc.add_paragraph()
    apply_num(p, num_b)
    set_run_font(p.add_run(text))
    return p


def add_number(doc, text, num_n):
    p = doc.add_paragraph()
    apply_num(p, num_n)
    set_run_font(p.add_run(text))
    return p


def add_para(doc, text='', bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, italic=italic)
    else:
        set_run_font(p.add_run(text), italic=italic)
    return p


def add_formula(doc, formula):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.right_indent = Inches(0.35)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(formula)
    set_run_font(r, name=PRIMARY_FONT, size=10.5, italic=True, color='434343')
    return p


def add_source(doc, author_year, title, url, note=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run(f'{author_year}. {title}. '), size=10)
    add_hyperlink(p, url, url)
    if note:
        set_run_font(p.add_run(f' {note}'), size=10, color='555555')
    return p


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles['Normal']
normal.font.name = PRIMARY_FONT
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal._element.rPr.rFonts.set(qn('w:ascii'), PRIMARY_FONT)
normal._element.rPr.rFonts.set(qn('w:hAnsi'), PRIMARY_FONT)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), CJK_FONT)
normal._element.rPr.rFonts.set(qn('w:hint'), 'eastAsia')
normal_lang = OxmlElement('w:lang')
normal_lang.set(qn('w:val'), 'en-US')
normal_lang.set(qn('w:eastAsia'), 'zh-CN')
normal._element.rPr.append(normal_lang)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

for style_name, size, color, before, after in [
    ('Heading 1', 20, '000000', 20, 6),
    ('Heading 2', 16, '000000', 18, 6),
    ('Heading 3', 14, '434343', 16, 4),
]:
    s = styles[style_name]
    s.font.name = PRIMARY_FONT
    s.font.size = Pt(size)
    s.font.bold = False
    s.font.color.rgb = RGBColor.from_string(color)
    s._element.rPr.rFonts.set(qn('w:ascii'), PRIMARY_FONT)
    s._element.rPr.rFonts.set(qn('w:hAnsi'), PRIMARY_FONT)
    s._element.rPr.rFonts.set(qn('w:eastAsia'), CJK_FONT)
    s._element.rPr.rFonts.set(qn('w:hint'), 'eastAsia')
    style_lang = OxmlElement('w:lang')
    style_lang.set(qn('w:val'), 'en-US')
    style_lang.set(qn('w:eastAsia'), 'zh-CN')
    s._element.rPr.append(style_lang)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

if 'Code Block' not in [s.name for s in styles]:
    code_style = styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
else:
    code_style = styles['Code Block']
code_style.font.name = PRIMARY_FONT
code_style.font.size = Pt(9)
code_style.font.color.rgb = RGBColor.from_string('434343')
code_style.paragraph_format.left_indent = Inches(0.35)
code_style.paragraph_format.right_indent = Inches(0.35)
code_style.paragraph_format.space_before = Pt(4)
code_style.paragraph_format.space_after = Pt(8)
code_style.paragraph_format.line_spacing = 1.05

num_b, num_n = add_numbering(doc)

# Title block (plain paragraph, never Word Title style)
title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
set_run_font(title.add_run('Simulated Patient 话语真实性评估框架'), size=26, bold=False)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
set_run_font(subtitle.add_run('基于心理学构念、临床量表与会谈行为研究的可实现方案'), size=13, color='555555')

add_para(doc, '适用系统：浏览器内实时语音 simulated patient；STT → LLM persona/affective state → TTS → reactive avatar。本文档聚焦患者话语和语音的心理真实性，不评价图像清晰度、唇形同步或身份漂移。')

doc.add_heading('执行摘要', level=1)
add_para(doc, '本框架评价四个核心维度：A. Persona and fact consistency；B. Clinical fidelity；C. Disclosure realism；F. Linguistic naturalness。心理学论文能够直接定义人格、症状、隐瞒/披露倾向、治疗联盟和部分语言表现，但通常不能直接给一条 LLM 回复打“真实性分”。因此，完整测量链应为：心理构念 → 隐藏状态 → 可观察行为 → 评价指标 → 人类效标验证。')
add_bullet(doc, 'A 评价患者是否保留病例事实，并表现出稳定的“情境—反应”人格模式，而不是机械地每轮说得一样。', num_b)
add_bullet(doc, 'B 评价对话能否让盲评者恢复出预设的症状画像，以及症状是否覆盖、矛盾或发生不合理跳变。', num_b)
add_bullet(doc, 'C 评价敏感信息是否随着信任、提问方式、污名和自主权担忧逐步披露，而非开场即全部倾倒。', num_b)
add_bullet(doc, 'F 评价文字和音频特征是否落在匹配真人患者的分布内，并保留会谈中的迟疑、修正、简短回答和非最优表达。', num_b)
add_para(doc, '建议第一阶段不要合并成一个总分。先报告各子指标、置信区间和失败样例；获得临床心理学家评分及真人/标准化患者语料后，再估计权重。')

doc.add_heading('1. 测量原则与单位', level=1)
doc.add_heading('1.1 四个分析层级', level=2)
for item in [
    'Utterance level：单条患者回复是否与事实、症状和当前披露状态一致。',
    'Episode level：一个主题或敏感事实从被触及到披露、回避或拒绝的完整过程。',
    'Session level：整场会谈的症状画像、语言分布和治疗联盟变化。',
    'Repeated-simulation level：同一 persona 在多次随机生成中的均值、方差和情境反应是否稳定。',
]:
    add_bullet(doc, item, num_b)
add_para(doc, '人格和语言自然度不能只在单句层面评价。按照人格状态分布理论，真实的人在不同时间会有明显波动；评价重点应是多次模拟后的分布及情境化反应。')

doc.add_heading('1.2 评价输出总览', level=2)
table = doc.add_table(rows=1, cols=4)
headers = ['维度', '核心问题', '主要可计算指标', '主要心理学 grounding']
for i, h in enumerate(headers):
    p = table.rows[0].cells[i].paragraphs[0]
    set_run_font(p.add_run(h), size=9.5, bold=True)
repeat_table_header(table.rows[0])
rows = [
    ('A', '还是同一个患者吗？', '事实准确率；情境签名；分布距离；跨 session 稳定性', 'CAPS；trait density distribution；SP performance'),
    ('B', '临床状态是否可识别且连贯？', '症状恢复度；矛盾率；覆盖率；无效状态跳变率', 'PHQ-9；BHS；C-SSRS；EMA'),
    ('C', '是否以心理上合理的方式披露？', '过早泄露率；披露延迟；支持敏感性；Brier/C-index', 'Self-Concealment；DDI；WAI；自杀披露研究'),
    ('F', '说话是否像匹配的真人患者？', '文本/声学分布距离；会谈行为比例；人类真实性评分', '抑郁语言；绝对化语言；语速、停顿、反应延迟'),
]
for row in rows:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        p = cells[i].paragraphs[0]
        set_run_font(p.add_run(value), size=9)
set_table_geometry(table, [650, 2200, 3380, 3130])

doc.add_heading('2. A — Persona and fact consistency', level=1)
doc.add_heading('2.1 指标定义', level=2)
add_para(doc, 'Persona consistency 包含两个不同问题：病例事实是否正确，以及患者在不同情境下是否呈现稳定但非机械的人格与情感反应。事实一致性是病例工程问题；人格一致性则应由人格心理学中的 person–situation interaction 来定义。')

doc.add_heading('2.2 心理学 grounding', level=2)
add_para(doc, 'Mischel 与 Shoda 的 Cognitive-Affective Personality System（CAPS）认为，人格稳定性主要体现为可重复的 if–then behavioral signature：当某类心理情境出现时，个体倾向于产生某类认知、情感和行为反应。Fleeson 的经验采样研究则显示，同一个人的即时人格状态变化很大，但状态分布的平均值和形状具有稳定的个体差异。因此，患者不应逐轮完全一致，而应具有稳定的情境反应方向和合理的波动范围。')
add_source(doc, 'Mischel, W., & Shoda, Y. (1995)', 'A cognitive-affective system theory of personality: Reconceptualizing situations, dispositions, dynamics, and invariance in personality structure', 'https://doi.org/10.1037/0033-295X.102.2.246')
add_source(doc, 'Fleeson, W. (2001)', 'Toward a structure- and process-integrated view of personality: Traits as density distributions of states', 'https://doi.org/10.1037/0022-3514.80.6.1011')
add_para(doc, '标准化患者研究提供了病例呈现的工程化测量范例：分别评价 script accuracy、general verbal communication、affective demeanor 和 realism，而不是用单一“一致性”分数代替。')
add_source(doc, 'Erby, L. A. H., Roter, D. L., & Biesecker, B. B. (2011)', 'Examination of standardized patient performance: Accuracy and consistency of six standardized patients over time', 'https://doi.org/10.1016/j.pec.2010.10.005')

doc.add_heading('2.3 A1 — Case fact accuracy', level=2)
add_para(doc, '定义：当某条病例事实被直接询问、被自然触发，或按照脚本本应主动出现时，患者是否以正确粒度表达该事实。每个事实必须同时保存 ground truth、敏感度、是否允许隐瞒以及触发条件。')
add_formula(doc, 'FactAccuracy = correct_fact_presentations / fact_opportunities')
add_para(doc, '建议对每次 opportunity 标记以下互斥状态：')
for item in [
    'correctly_disclosed：正确披露；',
    'correctly_withheld：依据 persona 和披露规则合理隐瞒；',
    'intentionally_denied：患者因害怕、污名或自主权担忧而否认；',
    'contradicted_by_error：没有心理原因、由模型遗忘导致的矛盾；',
    'not_asked / not_triggered：本次没有披露机会。',
]:
    add_bullet(doc, item, num_b)
add_para(doc, '计算时只有 contradicted_by_error 直接算事实错误。intentionally_denied 是否合理，应由 C 维度的披露模型评价，不能简单处罚。')

doc.add_heading('2.4 A2 — If–then behavioral signature', level=2)
add_para(doc, '构造一组固定的 interviewer probe，使它们只改变心理情境，不改变所问事实。例如：neutral question、empathic reflection、direct suicide question、judgmental response、premature problem solving、autonomy support。')
add_para(doc, '对每轮患者反应提取：disclosure、defensiveness、trust、anxiety、warmth、talkativeness、emotional intensity。可以由两个盲评心理学家评分，也可以在建立可靠人工标签后训练分类器。')
add_formula(doc, 'Delta(c,k) = E[state_k | context_c] - E[state_k | neutral]')
add_formula(doc, 'SignatureScore = SpearmanCorr(Delta_generated, Delta_target_or_human)')
add_para(doc, '如果 persona 规定“被评判时防御上升、披露下降”，而模型在该条件下反而主动详述敏感经历，则方向错误。Spearman correlation 适合早期实现，因为它优先评价反应方向和排序，不要求初始系数已经精确校准。')

doc.add_heading('2.5 A3 — 状态分布与跨 session 稳定性', level=2)
add_para(doc, '同一 persona、同一脚本至少生成 10–20 次。对每个心理状态计算 session mean、within-session variance 和条件下的变化。使用 Wasserstein distance 或 Maximum Mean Discrepancy 比较生成分布与真人/SP reference distribution。')
add_formula(doc, 'DistributionFidelity_k = exp(-WassersteinDistance(P_generated,k, P_reference,k))')
add_para(doc, '推荐同时报告：session-level mean 的 ICC、turn-level variability ratio，以及极端反应出现率。目标不是“ICC 越高越好”：过高且方差接近零可能意味着患者机械化；需要同时接近真人分布的均值和方差。')

doc.add_heading('2.6 实现所需数据', level=2)
for item in [
    '结构化病例事实表和每条事实的 disclosure policy；',
    '5–8 类标准情境 probe；',
    '每个 persona 至少 10 次随机复现；',
    '心理状态的人工行为锚点；',
    '最好有匹配真人或训练良好的 SP 作为分布参照。',
]:
    add_bullet(doc, item, num_b)

doc.add_heading('3. B — Clinical fidelity', level=1)
doc.add_heading('3.1 指标定义', level=2)
add_para(doc, 'Clinical fidelity 评价生成对话是否忠实表达预设的症状维度、严重程度和时间动态。它不应只问“像不像抑郁症”，而应评价具体症状能否被观察到、是否互相矛盾，以及是否随会谈事件合理变化。')

doc.add_heading('3.2 心理学与临床测量 grounding', level=2)
add_para(doc, '建议用经过验证的量表维度定义隐藏 clinical state，而不是让 LLM 自己维持一个自由文本诊断标签。PHQ-9 可定义抑郁症状频率；Beck Hopelessness Scale 可区分绝望与一般负面情绪；C-SSRS 将死亡愿望、主动自杀意念、方法、意图、计划和行为分开。')
add_source(doc, 'Kroenke, K., Spitzer, R. L., & Williams, J. B. W. (2001)', 'The PHQ-9: Validity of a brief depression severity measure', 'https://doi.org/10.1046/j.1525-1497.2001.016009606.x')
add_source(doc, 'Beck, A. T., Weissman, A., Lester, D., & Trexler, L. (1974)', 'The measurement of pessimism: The Hopelessness Scale', 'https://doi.org/10.1037/h0037562')
add_source(doc, 'Posner, K., et al. (2011)', 'The Columbia–Suicide Severity Rating Scale: Initial validity and internal consistency findings from three multisite studies', 'https://doi.org/10.1176/appi.ajp.2011.10111704')
add_para(doc, '自杀想法不是固定不变的标签。EMA 研究表明，自杀意念及其相关风险因素在小时级别可能发生显著波动。因此，评价应检查“变化是否有心理触发与时间尺度支持”，而不是要求全场恒定。')
add_source(doc, 'Kleiman, E. M., et al. (2017)', 'Examination of real-time fluctuations in suicidal ideation and its risk factors: Results from two ecological momentary assessment studies', 'https://doi.org/10.1037/abn0000273')

doc.add_heading('3.3 Clinical ground-truth schema', level=2)
code = doc.add_paragraph(style='Code Block')
code_text = '''clinical_state:\n  depression:\n    low_mood: 3\n    anhedonia: 3\n    sleep_problem: 2\n    fatigue: 3\n    guilt: 2\n    concentration: 1\n    psychomotor_retardation: 2\n  hopelessness: 0.80\n  suicidality:\n    wish_to_be_dead: true\n    active_ideation: true\n    method: true\n    intent: false\n    plan: false\n    past_behavior: true'''
set_run_font(code.add_run(code_text), size=9, color='434343')
add_para(doc, '量表用于定义模拟患者隐藏状态和评估维度，不代表可以仅凭一条生成语句诊断真实患者。不得把 C-SSRS 简化成单一“风险预测分”；它在本系统中主要用于保持意念和行为构成的一致性。')

doc.add_heading('3.4 B1 — Symptom coverage', level=2)
add_formula(doc, 'Coverage = symptoms_with_observable_evidence / symptoms_expected_to_be_observable')
add_para(doc, '只有在访谈问题为症状提供机会时才进入分母。患者不主动谈及睡眠问题，并不等于系统漏掉症状；如果治疗者明确问睡眠，而患者无心理理由地回答与设定相反，才属于覆盖或矛盾问题。')

doc.add_heading('3.5 B2 — Symptom recoverability', level=2)
add_para(doc, '将 persona 设定对评价者隐藏，让临床心理学家仅凭 transcript 或音视频估计症状维度和严重程度。')
add_formula(doc, 'Recovery_k = 1 - abs(estimated_state_k - target_state_k) / scale_range_k')
add_para(doc, '连续维度可报告 MAE、Spearman correlation 和 concordance correlation coefficient（CCC）；二元症状可报告 sensitivity、specificity 和 macro-F1。核心问题是：模拟患者有没有通过自然对话表达出预设的临床画像，而不是有没有复述量表条目。')

doc.add_heading('3.6 B3 — Clinical contradiction rate', level=2)
add_formula(doc, 'ContradictionRate = unsupported_clinical_contradictions / clinical_claims')
add_para(doc, '评价器需要区分三类现象：症状真实变化；因隐瞒而产生的表面否认；模型错误导致的矛盾。只有第三类直接进入 contradiction numerator。前两类分别由动态模型和 disclosure model 解释。')

doc.add_heading('3.7 B4 — Dynamic coherence', level=2)
add_formula(doc, 'InvalidTransitionRate = unsupported_state_jumps / all_state_transitions')
add_para(doc, '为每个状态定义最大变化速度和允许触发事件。例如，准确共情可在数轮内提高信任并降低紧张，但不应让重度快感缺失瞬间消失；提及近期失败可能短暂提高绝望；建立可接受的安全计划可能降低当下紧迫感，但不等于自杀史被改写。')

doc.add_heading('3.8 建议的 clinical 子指标', level=2)
for item in [
    'Symptom coverage：被适当询问的症状有多少获得可观察证据；',
    'Construct recoverability：盲评者能否恢复出目标症状画像；',
    'Contradiction rate：无心理解释的临床矛盾比例；',
    'Severity calibration：低、中、高严重度 persona 能否被正确排序；',
    'Dynamic coherence：状态变化是否符合触发事件和时间尺度；',
    'Suicide-structure accuracy：意念、方法、意图、计划与行为是否被正确区分。',
]:
    add_bullet(doc, item, num_b)

doc.add_heading('4. C — Disclosure realism', level=1)
doc.add_heading('4.1 指标定义', level=2)
add_para(doc, 'Disclosure realism 评价患者何时、向谁、在什么提问和关系条件下披露何种敏感信息。真实披露不是固定脚本顺序，而是由隐瞒倾向、信息敏感度、污名、预期后果、信任和治疗联盟共同影响的概率过程。')

doc.add_heading('4.2 心理学 grounding', level=2)
add_para(doc, 'Self-Concealment Scale 将 self-concealment 定义为主动隐藏被自己认为负面或痛苦的个人信息；它与简单的“话少”或一般自我披露并不等价。Distress Disclosure Index 则测量一个人倾向于隐瞒还是表达心理痛苦，并有研究将自报告倾向与结构化访谈中的实际承认、否认和可观察痛苦联系起来。')
add_source(doc, 'Larson, D. G., & Chastain, R. L. (1990)', 'Self-concealment: Conceptualization, measurement, and health implications', 'https://doi.org/10.1521/jscp.1990.9.4.439')
add_source(doc, 'Kahn, J. H., & Hessling, R. M. (2001)', 'Measuring the tendency to conceal versus disclose psychological distress', 'https://doi.org/10.1521/jscp.20.1.41.22254')
add_para(doc, '自杀披露研究提供了更具体的机制：污名担忧是准确披露的常见障碍；希望获得情感支持和被理解是促进因素。患者也可能担心对方过度反应、失去自主权或认为披露没有帮助。')
add_source(doc, 'Hom, M. A., Stanley, I. H., Podlogar, M. C., & Joiner, T. E. (2017)', '“Are you having thoughts of suicide?” Examining experiences with disclosing and denying suicidal ideation', 'https://doi.org/10.1002/jclp.22440')
add_source(doc, 'Richards, J. E., et al. (2019)', 'Understanding why patients may not report suicidal ideation at a health care visit prior to a suicide attempt: A qualitative study', 'https://doi.org/10.1176/appi.ps.201800342')
add_para(doc, 'Working Alliance Inventory 将治疗联盟组织为目标、任务和关系纽带，可作为信任状态和披露变化的关系性 grounding。')
add_source(doc, 'Horvath, A. O., & Greenberg, L. S. (1989)', 'Development and validation of the Working Alliance Inventory', 'https://doi.org/10.1037/0022-0167.36.2.223')

doc.add_heading('4.3 隐藏事实的数据结构', level=2)
code = doc.add_paragraph(style='Code Block')
code_text = '''hidden_fact:\n  id: overdose_consideration\n  ground_truth: true\n  sensitivity: 0.90\n  spontaneous_disclosure_threshold: 0.88\n  direct_question_threshold: 0.62\n  fear_of_consequence: 0.80\n  may_initially_deny: true\n  corrective_conditions:\n    - nonjudgmental_follow_up\n    - autonomy_support\n    - accurate_reflection'''
set_run_font(code.add_run(code_text), size=9, color='434343')

doc.add_heading('4.4 Trust/alliance 状态更新', level=2)
add_formula(doc, 'T(t+1) = clip(rho*T(t) + b1*Empathy + b2*AccurateReflection + b3*Validation + b4*AutonomySupport - b5*Judgment - b6*PrematureAdvice - b7*Coercion)')
add_para(doc, '初期没有人类数据时，不要声称这些 beta 系数是心理学论文直接给出的。先用文献支持的单调约束：共情和准确理解不应系统性降低披露；污名、判断和失去自主权的担忧不应系统性提高披露。之后用真人会谈或 SP 实验拟合系数。')

doc.add_heading('4.5 披露概率', level=2)
add_formula(doc, 'P(disclose_i,t) = sigmoid(a + wT*T(t) + wQ*QuestionDirectness + wA*Alliance - wS*Sensitivity_i - wC*SelfConcealment - wG*Stigma - wL*FearOfLostAutonomy)')
add_para(doc, '输出不应只有 disclose / not disclose。更真实的中间行为包括：模糊回答、最小化、改变时间范围、承认情绪但否认意图、反问后果、沉默、要求澄清、先否认后修正。')

doc.add_heading('4.6 具体指标与算法', level=2)
table = doc.add_table(rows=1, cols=4)
for i, h in enumerate(['子指标', '计算方法', '分析单位', '失败含义']):
    set_run_font(table.rows[0].cells[i].paragraphs[0].add_run(h), size=9.5, bold=True)
repeat_table_header(table.rows[0])
metric_rows = [
    ('Early leakage rate', '建立足够信任或被直接询问前出现的高敏感事实 / 高敏感事实总数', 'session', '过度配合、信息倾倒'),
    ('Time to disclosure', '从主题首次被触及到首次有效披露的 turn 数；可做 survival analysis', 'episode', '披露过快或永久卡住'),
    ('Disclosure depth curve', '每轮累计披露深度；与真人/SP 曲线比较', 'session', '缺少渐进性'),
    ('Support sensitivity', '支持性条件的披露概率 - 中性/评判条件的披露概率', 'counterfactual pair', '对关系情境不敏感'),
    ('Question contingency', '直接提问后出现回答、回避或明确拒绝的比例', 'turn pair', '忽略问题或随机换题'),
    ('Brier score', '披露概率预测与实际披露事件的均方误差', 'fact-turn', '概率未校准'),
    ('C-index', '高敏感信息是否比低敏感信息更晚披露', 'session', '披露顺序与敏感度无关'),
]
for row in metric_rows:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        set_run_font(cells[i].paragraphs[0].add_run(value), size=8.6)
set_table_geometry(table, [1750, 3670, 1650, 2290])

doc.add_heading('4.7 最关键的 counterfactual test', level=2)
add_para(doc, '同一 persona、同一问题和同一随机种子，仅改变治疗者回应方式：supportive、neutral、judgmental/coercive。比较后续 3–5 轮的披露概率、深度、防御和信任。如果模型的披露不随这些条件发生方向一致的变化，即使单句很流畅，心理过程仍不真实。')

doc.add_heading('5. F — Linguistic naturalness', level=1)
doc.add_heading('5.1 指标定义', level=2)
add_para(doc, 'Linguistic naturalness 应拆成四部分：语义表达是否符合心理状态、会谈行为是否像患者、文本风格是否落在匹配人群分布内、声音时间特征是否符合症状和情境。心理学能够 grounding 症状相关语言特征，但不能提供一个对所有患者通用的“自然度量表”。最终必须使用匹配的真人/SP 会谈语料。')

doc.add_heading('5.2 心理语言学 grounding', level=2)
add_para(doc, '抑郁相关研究发现，当前抑郁参与者的书面语言中第一人称单数和负性情绪词更多。另一项网络论坛研究发现，焦虑、抑郁和自杀意念语料中绝对化词语更多，自杀意念论坛尤其明显。由于这些研究包含书面文章或论坛文本，它们只能提供特征方向，不能直接作为临床口语阈值。')
add_source(doc, 'Rude, S. S., Gortner, E.-M., & Pennebaker, J. W. (2004)', 'Language use of depressed and depression-vulnerable college students', 'https://doi.org/10.1080/02699930441000030')
add_source(doc, 'Al-Mosaiwi, M., & Johnstone, T. (2018)', 'In an absolute state: Elevated use of absolutist words is a marker specific to anxiety, depression, and suicidal ideation', 'https://doi.org/10.1177/2167702617747074')
add_para(doc, '临床访谈语音研究更适合你的系统。一项包含 241 名参与者和 1,058 份数据的研究测量了 speech rate、pause time 和 response time，并考察其与抑郁严重程度的关系。较早的声学研究也发现抑郁组问题后反应延迟更长。')
add_source(doc, 'Yamamoto, M., et al. (2020)', 'Using speech recognition technology to investigate the association between timing-related speech features and depression severity', 'https://doi.org/10.1371/journal.pone.0238726')
add_source(doc, 'Nilsonne, A. (1988)', 'Speech characteristics as indicators of depressive illness', 'https://doi.org/10.1111/j.1600-0447.1988.tb05118.x')

doc.add_heading('5.3 文本特征', level=2)
for item in [
    'first_person_singular_rate；negative/positive emotion rate；absolutist word rate；',
    'hedging、模糊词、最小化表达和 uncertainty markers；',
    'mean utterance length、type-token ratio、重复率、未完成句比例；',
    '自我修正、澄清请求、沉默后的简短回答、拒绝回答；',
    '患者 talk / clinician talk ratio，以及 clinical、psychosocial、emotional 和 facilitative talk 的比例。',
]:
    add_bullet(doc, item, num_b)
add_para(doc, '最后一组会谈行为比例可参考标准化患者研究中对 RIAS 类别和 verbal activity 的使用，但目标分布必须按病例和访谈阶段分层；不存在“患者说得越多越自然”的统一规律。')

doc.add_heading('5.4 音频特征', level=2)
for item in [
    'speech rate：每分钟词数或 syllables per second；',
    'response latency：治疗者结束到患者开始说话的时间；',
    'within-turn pause：患者发言内部停顿的均值、分位数和总比例；',
    'pitch mean / variance、energy mean / variance；',
    'voice activity ratio、cut-off/repair rate、被打断后的恢复方式。',
]:
    add_bullet(doc, item, num_b)
add_para(doc, '这些特征应直接从 Deepgram 时间戳、VAD、原始音频和 Cartesia 输出中提取，而不是只从 transcript 推断。TTS 本身会限制停顿、语调和 disfluency，因此 F 维度需分别报告 text naturalness 与 rendered-speech naturalness。')

doc.add_heading('5.5 分布式计算', level=2)
add_formula(doc, 'FeatureVector = [lexical, dialogue-act, timing, acoustic features]')
add_formula(doc, 'NaturalnessFidelity = exp(-Distance(X_generated, X_matched_human))')
add_para(doc, 'Distance 可用 Mahalanobis distance、Wasserstein distance 或 MMD。Reference corpus 至少要匹配诊断/症状严重度、年龄、语言、访谈阶段、问题类型、临床评估 versus 心理治疗，以及音频采集方式。否则指标可能测到年龄、任务或麦克风差异。')

doc.add_heading('5.6 必须避免的实现错误', level=2)
for item in [
    '不要最大化“抑郁语言标记”；目标是接近分布，不是负性词越多越好。',
    '不要把论坛文字的 effect size 直接当作临床口语阈值。',
    '不要只用 perplexity 或 LLM-as-judge 代替患者语言行为。',
    '不要把语法完美当作自然；真实患者会迟疑、修正、答非所问或只回答一部分。',
    '不要将语言标记作为真实患者诊断工具；这里它们仅用于模拟行为评价。',
]:
    add_bullet(doc, item, num_b)

doc.add_heading('6. 人类评价与心理测量验证', level=1)
doc.add_heading('6.1 Blinded clinician rating', level=2)
add_para(doc, '至少邀请两名不了解 persona 条件的临床心理学家独立评分。连续评分报告 ICC；类别标签报告 Krippendorff’s alpha 或 weighted kappa。出现分歧时保留原始评分，并通过第三位评价者形成 adjudicated label，不能只保留共识结果。')
add_para(doc, '建议 1–5 分锚点：')
for item in [
    '1：明显错误或心理过程不可能；',
    '2：存在多处矛盾、过度披露或刻板化表达；',
    '3：基本可接受，但仍可辨认出系统化模式；',
    '4：大部分符合临床经验，仅有轻微问题；',
    '5：在该情境下与真实患者/SP 难以区分。',
]:
    add_bullet(doc, item, num_b)

doc.add_heading('6.2 Construct validity', level=2)
add_para(doc, '心理学量表只提供 construct definition；你的 transcript metric 仍需要证明构念效度。建议依次验证：')
for item in [
    'Known-groups validity：高 versus 低 hopelessness、self-concealment 或初始信任 persona 是否被指标区分；',
    'Convergent validity：自动指标是否与临床心理学家评分、量表反推值相关；',
    'Discriminant validity：披露分数不应只是抑郁严重度或回复长度的替代变量；',
    'Criterion validity：生成患者与真实患者/SP 在匹配任务中的分布距离；',
    'Generalizability：在未见过的 persona、访谈者表达和随机种子上是否保持。',
]:
    add_number(doc, item, num_n)

doc.add_heading('7. 推荐实验设计', level=1)
doc.add_heading('7.1 Persona sampling', level=2)
add_para(doc, '先建立 16–32 个结构化 persona，至少正交改变以下变量：抑郁严重度、绝望、自杀意念结构、self-concealment、stigma/fear of consequences、初始信任和人际防御。避免让所有高抑郁 persona 都同时高隐瞒，否则无法区分指标。')

doc.add_heading('7.2 标准化 interviewer conditions', level=2)
for item in [
    'Supportive：准确反映、验证体验、尊重自主权；',
    'Neutral：标准临床提问，不额外强化关系；',
    'Premature problem solving：过早建议解决方案；',
    'Judgmental/coercive：表现担忧但带有判断或控制暗示；',
    'Direct suicide assessment：直接、清晰询问意念、方法、意图和计划。',
]:
    add_bullet(doc, item, num_b)

doc.add_heading('7.3 重复与样本单位', level=2)
add_para(doc, '每个 persona × interviewer condition 至少运行 10 次，保留正常生产温度和完整音频。分析时把 persona 和访谈脚本作为随机效应，避免把几百条 turn 错当成几百个独立患者。')

doc.add_heading('7.4 推荐统计模型', level=2)
for item in [
    '连续心理状态：mixed-effects regression；random intercepts for persona and script。',
    '是否披露：mixed-effects logistic regression。',
    '披露时间：survival model / discrete-time hazard model。',
    '重复 session 的分布：Wasserstein/MMD + bootstrap confidence interval。',
    '评价者一致性：ICC 或 Krippendorff’s alpha。',
]:
    add_bullet(doc, item, num_b)

doc.add_heading('8. 数据结构建议', level=1)
code = doc.add_paragraph(style='Code Block')
code_text = '''persona_id: P017\ntraits:\n  initial_trust: 0.35\n  defensiveness: 0.70\nclinical_state:\n  depression_severity: 0.78\n  hopelessness: 0.82\n  suicidality:\n    active_ideation: true\n    method: true\n    intent: false\ndisclosure:\n  self_concealment: 0.76\n  stigma: 0.68\n  fear_of_lost_autonomy: 0.84\nif_then_rules:\n  empathic_reflection:\n    trust_delta: [0.04, 0.10]\n    disclosure_delta: [0.03, 0.12]\n  judgment:\n    defensiveness_delta: [0.10, 0.25]\n    disclosure_delta: [-0.20, -0.05]\nhidden_facts:\n  - fact_id: overdose_consideration\n    sensitivity: 0.90\n    may_initially_deny: true\nspeech_profile:\n  target_speech_rate_z: -0.7\n  target_response_latency_z: 0.8'''
set_run_font(code.add_run(code_text), size=9, color='434343')

doc.add_heading('9. Evaluation pipeline', level=1)
for item in [
    '冻结 persona ground truth 和标准化 interviewer script；',
    '对每个条件运行多次完整语音 session，并保存 transcript、word timestamps、VAD、音频和内部状态；',
    '对 utterance 做 fact、clinical evidence、disclosure act 和 dialogue act 标注；',
    '从 transcript/audio 提取语言和声学特征；',
    '计算 A/B/C/F 子指标和 bootstrap confidence interval；',
    '把盲评 clinician score 作为外部效标，检验自动指标的收敛和区分效度；',
    '输出维度级 dashboard、失败案例和 counterfactual 对照，不急于给单一总分。',
]:
    add_number(doc, item, num_n)

doc.add_heading('10. MVP：最先实现的指标', level=1)
add_para(doc, '如果资源有限，建议先实现下列八项。它们覆盖四个维度，并且不依赖大型真人语料才能开始：')
for item in [
    'A1 Fact accuracy / unsupported contradiction；',
    'A2 六类情境 probe 的 reaction-direction pass rate；',
    'B1 Symptom coverage；',
    'B2 Clinician-based symptom recoverability；',
    'C1 Early leakage rate；',
    'C2 Support sensitivity 的 counterfactual test；',
    'F1 speech rate、response latency、pause ratio；',
    'F2 blinded clinician realism rating。',
]:
    add_bullet(doc, item, num_b)
add_para(doc, '第二阶段再加入 disclosure survival model、分布距离、自动行为分类器和权重学习。这样可以先验证 construct 是否成立，再提高自动化程度。')

doc.add_heading('11. 推荐报告模板', level=1)
table = doc.add_table(rows=1, cols=5)
for i, h in enumerate(['Dimension', 'Metric', 'Estimate', '95% CI', 'Human criterion / notes']):
    set_run_font(table.rows[0].cells[i].paragraphs[0].add_run(h), size=9.2, bold=True)
repeat_table_header(table.rows[0])
report_rows = [
    ('A', 'Fact accuracy', '0.92', '[0.88, 0.95]', '区分 intentional denial'),
    ('A', 'Signature correlation', '0.71', '[0.59, 0.80]', '目标为情境方向'),
    ('B', 'Symptom recovery CCC', '0.68', '[0.54, 0.77]', '盲评 clinician'),
    ('C', 'Early leakage rate', '0.11', '[0.07, 0.16]', '高敏感事实'),
    ('C', 'Support sensitivity', '+0.24', '[0.16, 0.31]', 'supportive - neutral'),
    ('F', 'Human realism', '3.8/5', '[3.5, 4.1]', '与 matched SP 比较'),
]
for row in report_rows:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        set_run_font(cells[i].paragraphs[0].add_run(value), size=8.7)
set_table_geometry(table, [800, 2200, 1200, 1350, 3810])
add_para(doc, '表中数值仅展示报告格式，不是推荐阈值。阈值必须在 pilot 中依据真人/SP 分布、评价者一致性和任务用途确定。高风险考试用途需要比形成性训练更严格的一致性要求。')

doc.add_heading('12. 重要限制与伦理注意', level=1)
for item in [
    '心理量表定义的是构念，不会自动验证生成文本。必须另外验证从构念到行为的映射。',
    '部分量表或完整条目可能受版权或授权限制；实现前确认使用许可。本文档未复制量表条目。',
    '语言标记只能作为模拟真实性特征，不能用于真实患者的单独诊断或自杀预测。',
    '不要把“更严重、更负面、更迟缓”误认为“更真实”；目标是匹配设定与参考分布。',
    '若 simulated patient 用于评价学员，患者自身的不一致会成为测量误差，需要用多 persona、多情境和多次复现控制。',
    '自杀主题脚本应由有相关临床经验的心理学家审核，并明确系统是教学模拟而非临床服务。',
]:
    add_bullet(doc, item, num_b)

doc.add_heading('13. Grounding paper 对照表', level=1)
table = doc.add_table(rows=1, cols=4)
for i, h in enumerate(['指标/构念', '核心论文', '论文直接提供什么', '本项目仍需验证什么']):
    set_run_font(table.rows[0].cells[i].paragraphs[0].add_run(h), size=9.2, bold=True)
repeat_table_header(table.rows[0])
ground_rows = [
    ('A 情境化人格', 'Mischel & Shoda, 1995', 'if–then behavioral signature 理论', '哪些 probe 与行为维度适合本 persona'),
    ('A 状态分布', 'Fleeson, 2001', '均值稳定、即时状态高波动', '真人/SP 的目标均值和方差'),
    ('A 病例准确性', 'Erby et al., 2011', 'accuracy、communication、affect、realism 分开测', 'LLM 中 intentional denial 的编码'),
    ('B 抑郁症状', 'Kroenke et al., 2001', 'PHQ-9 症状严重度构念', '症状如何自然体现在会谈中'),
    ('B 自杀结构', 'Posner et al., 2011', '意念、方法、意图、计划、行为的区分', '披露/否认机制及动态变化'),
    ('B 动态状态', 'Kleiman et al., 2017', '短时间内自杀意念与相关因素会波动', '单场访谈的转移参数'),
    ('C 隐瞒倾向', 'Larson & Chastain, 1990', 'self-concealment 构念与量表', '从量表倾向到 turn-level 行为'),
    ('C 痛苦披露', 'Kahn & Hessling, 2001', 'conceal–disclose 连续维度', '不同信息敏感度和提问条件'),
    ('C 自杀披露', 'Hom et al., 2017; Richards et al., 2019', '污名、支持、理解、后果与自主权因素', '这些因素在模拟会谈中的效应大小'),
    ('C 治疗联盟', 'Horvath & Greenberg, 1989', 'goal、task、bond 的联盟结构', 'turn-level trust 更新模型'),
    ('F 文本语言', 'Rude et al., 2004; Al-Mosaiwi & Johnstone, 2018', '若干抑郁/自杀相关语言方向', '临床口语中的基线与阈值'),
    ('F 语音时间', 'Yamamoto et al., 2020; Nilsonne, 1988', 'speech rate、pause、response latency', '当前 TTS 与匹配人群的分布'),
]
for row in ground_rows:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        set_run_font(cells[i].paragraphs[0].add_run(value), size=8.2)
set_table_geometry(table, [1800, 2200, 2620, 2740])

doc.add_heading('14. References', level=1)
refs = [
    ('Al-Mosaiwi, M., & Johnstone, T. (2018)', 'In an absolute state: Elevated use of absolutist words is a marker specific to anxiety, depression, and suicidal ideation. Clinical Psychological Science, 6(4), 529–542.', 'https://doi.org/10.1177/2167702617747074'),
    ('Beck, A. T., Weissman, A., Lester, D., & Trexler, L. (1974)', 'The measurement of pessimism: The Hopelessness Scale. Journal of Consulting and Clinical Psychology, 42(6), 861–865.', 'https://doi.org/10.1037/h0037562'),
    ('Erby, L. A. H., Roter, D. L., & Biesecker, B. B. (2011)', 'Examination of standardized patient performance: Accuracy and consistency of six standardized patients over time. Patient Education and Counseling, 85(2), 194–200.', 'https://doi.org/10.1016/j.pec.2010.10.005'),
    ('Fleeson, W. (2001)', 'Toward a structure- and process-integrated view of personality: Traits as density distributions of states. Journal of Personality and Social Psychology, 80(6), 1011–1027.', 'https://doi.org/10.1037/0022-3514.80.6.1011'),
    ('Hom, M. A., Stanley, I. H., Podlogar, M. C., & Joiner, T. E. (2017)', '“Are you having thoughts of suicide?” Examining experiences with disclosing and denying suicidal ideation. Journal of Clinical Psychology, 73(10), 1382–1392.', 'https://doi.org/10.1002/jclp.22440'),
    ('Horvath, A. O., & Greenberg, L. S. (1989)', 'Development and validation of the Working Alliance Inventory. Journal of Counseling Psychology, 36(2), 223–233.', 'https://doi.org/10.1037/0022-0167.36.2.223'),
    ('Kahn, J. H., & Hessling, R. M. (2001)', 'Measuring the tendency to conceal versus disclose psychological distress. Journal of Social and Clinical Psychology, 20(1), 41–65.', 'https://doi.org/10.1521/jscp.20.1.41.22254'),
    ('Kleiman, E. M., Turner, B. J., Fedor, S., Beale, E. E., Huffman, J. C., & Nock, M. K. (2017)', 'Examination of real-time fluctuations in suicidal ideation and its risk factors: Results from two ecological momentary assessment studies. Journal of Abnormal Psychology, 126(6), 726–738.', 'https://doi.org/10.1037/abn0000273'),
    ('Kroenke, K., Spitzer, R. L., & Williams, J. B. W. (2001)', 'The PHQ-9: Validity of a brief depression severity measure. Journal of General Internal Medicine, 16(9), 606–613.', 'https://doi.org/10.1046/j.1525-1497.2001.016009606.x'),
    ('Larson, D. G., & Chastain, R. L. (1990)', 'Self-concealment: Conceptualization, measurement, and health implications. Journal of Social and Clinical Psychology, 9(4), 439–455.', 'https://doi.org/10.1521/jscp.1990.9.4.439'),
    ('Mischel, W., & Shoda, Y. (1995)', 'A cognitive-affective system theory of personality: Reconceptualizing situations, dispositions, dynamics, and invariance in personality structure. Psychological Review, 102(2), 246–268.', 'https://doi.org/10.1037/0033-295X.102.2.246'),
    ('Nilsonne, A. (1988)', 'Speech characteristics as indicators of depressive illness. Acta Psychiatrica Scandinavica, 77(3), 253–263.', 'https://doi.org/10.1111/j.1600-0447.1988.tb05118.x'),
    ('Posner, K., et al. (2011)', 'The Columbia–Suicide Severity Rating Scale: Initial validity and internal consistency findings from three multisite studies with adolescents and adults. American Journal of Psychiatry, 168(12), 1266–1277.', 'https://doi.org/10.1176/appi.ajp.2011.10111704'),
    ('Richards, J. E., et al. (2019)', 'Understanding why patients may not report suicidal ideation at a health care visit prior to a suicide attempt: A qualitative study. Psychiatric Services, 70(1), 40–45.', 'https://doi.org/10.1176/appi.ps.201800342'),
    ('Rude, S. S., Gortner, E.-M., & Pennebaker, J. W. (2004)', 'Language use of depressed and depression-vulnerable college students. Cognition and Emotion, 18(8), 1121–1133.', 'https://doi.org/10.1080/02699930441000030'),
    ('Yamamoto, M., et al. (2020)', 'Using speech recognition technology to investigate the association between timing-related speech features and depression severity. PLOS ONE, 15(9), e0238726.', 'https://doi.org/10.1371/journal.pone.0238726'),
]
for a, t, u in refs:
    add_source(doc, a, t, u)

# Document metadata and final paragraph control
doc.core_properties.title = 'Simulated Patient 话语真实性评估框架'
doc.core_properties.subject = 'Psychology-grounded evaluation metrics for simulated patient utterances'
doc.core_properties.author = 'Project working document'
doc.core_properties.keywords = 'simulated patient, evaluation, psychology, clinical fidelity, disclosure, language'

for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True

doc.save(OUT)
print(OUT)
